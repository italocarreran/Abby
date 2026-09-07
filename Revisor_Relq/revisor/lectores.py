# -*- coding: utf-8 -*-
"""Adaptadores de lectura Excel usados por los verificadores del Revisor.

Este módulo no conoce la ventana ni las reglas V4…V17. Conserva los fallbacks
openpyxl/xlwings y los diagnósticos OOXML del punto de entrada histórico.
"""

from pathlib import Path
import re
import sys

from __comun__ import excel_xml as _excel_xml
from __comun__.texto import suave_textual as normalizar
from revisor.archivos import mtime


def _suma_rango_openpyxl(ws, ref):
    total = 0.0
    hubo = False
    for fila in ws[ref]:
        celdas = fila if isinstance(fila, tuple) else (fila,)
        for c in celdas:
            if isinstance(c.value, (int, float)):
                total += float(c.value)
                hubo = True
    return total if hubo else None


CACHE_COLUMNAS = {}     # (ruta, mtime, hoja, col, col_filtro, fila) -> {clave: [suma, n]}


def col_letra(n):
    """4 -> "D".  Al reves de col_letra_a_num."""
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def col_letra_a_num(letra):
    n = 0
    for c in str(letra).upper():
        if "A" <= c <= "Z":
            n = n * 26 + (ord(c) - ord("A") + 1)
    return n


def _es_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def leer_columna_excel(ruta, hoja, columna, fila_inicio, col_filtro, valores_filtro, log):
    """Suma una columna completa desde fila_inicio hacia abajo.
    Si se indica col_filtro, suma solo las filas cuyo valor de esa columna esta
    en valores_filtro (comparacion sin tildes ni mayusculas).
    Devuelve (suma, n_filas, {valor_filtro: suma}) o (None, 0, {})."""
    ci = col_letra_a_num(columna)
    cf = col_letra_a_num(col_filtro) if col_filtro else 0
    objetivo = {normalizar(v) for v in (valores_filtro or [])}

    # Cache: la misma columna se lee una sola vez por corrida aunque varios
    # valores la usen con distinto filtro (p.ej. SCMT y SCPC).
    ck = (str(ruta), mtime(ruta), hoja, ci, cf, int(fila_inicio))

    def desde_desglose(bruto):
        """bruto: {clave: [suma, n]} -> (total, n, {clave: suma})"""
        if cf and objetivo:
            sel = [v for k, v in bruto.items() if k in objetivo]
        else:
            sel = list(bruto.values())
        total = sum(v[0] for v in sel)
        n = sum(v[1] for v in sel)
        return total, n, {k: v[0] for k, v in bruto.items()}

    if ck in CACHE_COLUMNAS:
        return desde_desglose(CACHE_COLUMNAS[ck])

    def acumular(pares):
        """pares: iterable de (monto, tipo)"""
        bruto = {}
        for monto, tipo in pares:
            if not _es_num(monto):
                continue
            clave = "(todo)" if not cf else normalizar(tipo)
            reg = bruto.setdefault(clave, [0.0, 0])
            reg[0] += float(monto)
            reg[1] += 1
        CACHE_COLUMNAS[ck] = bruto
        return desde_desglose(bruto)

    # --- openpyxl (valores cacheados) --------------------------------------
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(ruta), data_only=True, read_only=True)
        try:
            real = resolver_hoja(wb.sheetnames, hoja)
            if real is None:
                log(f"    ! La hoja '{hoja}' no existe. Hojas: {', '.join(wb.sheetnames)}")
                return None, 0, {}
            ws = wb[real]
            ancho = max(ci, cf)

            def gen():
                for fila in ws.iter_rows(min_row=fila_inicio, max_col=ancho,
                                         values_only=True):
                    m = fila[ci - 1] if len(fila) >= ci else None
                    t = fila[cf - 1] if cf and len(fila) >= cf else None
                    yield m, t

            total, n, desg = acumular(gen())
            if desg:
                return total, n, desg
        finally:
            wb.close()
        log("    · sin valores cacheados, reintentando con Excel...")
    except Exception as e:
        log(f"    · openpyxl no pudo leer ({e}); reintentando con Excel...")

    # --- xlwings ----------------------------------------------------------
    app = wb2 = None
    try:
        import xlwings as xw
        app = xw.App(visible=False, add_book=False)
        app.display_alerts = False
        app.screen_updating = False
        wb2 = app.books.open(str(ruta), read_only=True, update_links=False)
        nombres = [s.name for s in wb2.sheets]
        real = resolver_hoja(nombres, hoja)
        if real is None:
            log(f"    ! La hoja '{hoja}' no existe. Hojas: {', '.join(nombres)}")
            return None, 0, {}
        sh = wb2.sheets[real]
        ult = sh.used_range.last_cell.row
        if ult < fila_inicio:
            return 0.0, 0, {}
        montos = sh.range((fila_inicio, ci), (ult, ci)).value
        if not isinstance(montos, list):
            montos = [montos]
        if cf:
            tipos = sh.range((fila_inicio, cf), (ult, cf)).value
            if not isinstance(tipos, list):
                tipos = [tipos]
        else:
            tipos = [None] * len(montos)
        return acumular(zip(montos, tipos))
    except Exception as e:
        log(f"    ! Error leyendo con Excel: {e}")
        return None, 0, {}
    finally:
        try:
            if wb2 is not None:
                wb2.close()
            if app is not None:
                app.quit()
        except Exception:
            pass


NS_XL = _excel_xml.NS_XL
NS_REL = _excel_xml.NS_REL


es_zip_excel = _excel_xml.es_zip_excel


ubicar_hoja_xml = _excel_xml.ubicar_hoja_xml


expandir_columnas = _excel_xml.expandir_columnas


def buscar_marcas_rapido(ruta, hoja, fila_inicio, reglas, log, tope_detalle=30):
    """Busca errores de fórmula y textos prohibidos en columnas puntuales,
    escaneando el XML de la hoja por trozos en vez de cargar el libro.

    reglas: [{"rangos": ["CF:CI"], "errores": True, "textos": ["REVISAR"]}, ...]
    Devuelve {"conteo": {motivo: n}, "marcas": [(celda, motivo, valor)]} o None."""
    import zipfile
    import xml.etree.ElementTree as ET

    if not es_zip_excel(ruta):
        return None

    # columna -> indice de regla
    de_columna = {}
    for i, regla in enumerate(reglas or []):
        for r in regla.get("rangos", []):
            for letra in expandir_columnas(r):
                de_columna[letra.encode()] = i
    if not de_columna:
        return {"conteo": {}, "marcas": []}

    alternativas = b"|".join(sorted(de_columna, key=len, reverse=True))
    patron = re.compile(rb'<c r="(' + alternativas + rb')(\d+)"([^>]*?)(?:/>|>(.*?)</c>)',
                        re.S)
    fila_inicio = int(fila_inicio)
    conteo, marcas = {}, []

    def anotar(celda, motivo, valor):
        conteo[motivo] = conteo.get(motivo, 0) + 1
        if len(marcas) < tope_detalle:
            marcas.append((celda, motivo, str(valor)[:40]))

    try:
        with zipfile.ZipFile(str(ruta)) as z:
            ruta_hoja, hojas = ubicar_hoja_xml(z, hoja)
            if ruta_hoja is None:
                log(f"    ! La hoja '{hoja}' no existe. Hojas: {', '.join(hojas)}")
                return None

            # cadenas compartidas: se resuelven una vez y solo si hacen falta
            compartidas, malos = [], {}
            if any(regla.get("textos") for regla in reglas):
                if "xl/sharedStrings.xml" in set(z.namelist()):
                    ss = ET.fromstring(z.read("xl/sharedStrings.xml"))
                    compartidas = ["".join(si.itertext()) for si in ss.iter(f"{NS_XL}si")]
                # Una misma cadena puede estar prohibida por varias reglas, asi
                # que se guarda POR REGLA: {indice: {n_regla: texto}}
                for i, texto in enumerate(compartidas):
                    tn = normalizar(texto)
                    for idx, regla in enumerate(reglas):
                        for t in regla.get("textos", []):
                            if normalizar(t) in tn:
                                malos.setdefault(str(i).encode(), {})[idx] = t
                                break

            with z.open(ruta_hoja) as f:
                cola = b""
                while True:
                    trozo = f.read(1 << 20)
                    if not trozo:
                        break
                    buf = cola + trozo
                    fin = 0
                    for m in patron.finditer(buf):
                        fin = m.end()
                        col, nfila, attrs, cuerpo = m.group(1), m.group(2), m.group(3), m.group(4) or b""
                        if int(nfila) < fila_inicio:
                            continue
                        regla = reglas[de_columna[col]]
                        celda = (col + nfila).decode()
                        if b't="e"' in attrs:
                            if regla.get("errores"):
                                mv = re.search(rb"<v>([^<]*)</v>", cuerpo)
                                anotar(celda, "valor de error",
                                       mv.group(1).decode() if mv else "#?")
                            continue
                        if not regla.get("textos"):
                            continue
                        if b't="s"' in attrs:
                            mv = re.search(rb"<v>([^<]*)</v>", cuerpo)
                            porregla = malos.get(mv.group(1)) if mv else None
                            if porregla:
                                t = porregla.get(de_columna[col])
                                if t:
                                    anotar(celda, f"texto '{t}'",
                                           compartidas[int(mv.group(1))])
                        elif b'inlineStr' in attrs or b't="str"' in attrs:
                            # OJO: hay que leer SOLO el resultado, nunca el nodo
                            # <f> de la formula. Muchas formulas llevan la
                            # palabra buscada adentro (IFERROR(...,"REVISAR"))
                            # y mirar la formula marcaba toda la columna.
                            if b'inlineStr' in attrs:
                                mv = re.search(rb"<is>(.*?)</is>", cuerpo, re.S)
                                bruto = re.sub(rb"<[^>]+>", b"", mv.group(1)) if mv else b""
                            else:
                                mv = re.search(rb"<v>(.*?)</v>", cuerpo, re.S)
                                bruto = mv.group(1) if mv else b""
                            if not bruto:
                                continue
                            texto = bruto.decode("utf-8", "ignore")
                            texto = (texto.replace("&amp;", "&").replace("&lt;", "<")
                                     .replace("&gt;", ">").replace("&quot;", '"')
                                     .replace("&apos;", "'"))
                            tn = normalizar(texto)
                            for t in regla["textos"]:
                                if normalizar(t) in tn:
                                    anotar(celda, f"texto '{t}'", texto.strip())
                                    break
                    cola = buf[fin:] if fin else buf[-8192:]
    except Exception as e:
        log(f"    ! No se pudo escanear {Path(ruta).name}: {e}")
        return None
    return {"conteo": conteo, "marcas": marcas}


def leer_celdas_rapido(ruta, hoja, celdas):
    """Lee celdas puntuales de un .xlsx/.xlsm sin cargar el libro completo.

    Un .xlsm es un ZIP con XML adentro. Se recorre el XML de la hoja en
    streaming y se corta en cuanto se pasa de la ultima fila pedida, asi que
    para celdas de las primeras filas (H1, EE6) casi no se lee nada, aunque el
    archivo tenga miles de filas y millones de formulas.

    Devuelve {celda: valor} con los valores YA CALCULADOS que Excel dejo
    guardados. Si el archivo nunca fue calculado y guardado, no habra valores.
    Devuelve None si no se pudo (formato distinto, hoja inexistente, etc.)."""
    import zipfile
    import xml.etree.ElementTree as ET

    if not es_zip_excel(ruta):
        return None

    pedidas = {c.upper().replace("$", "") for c in celdas}
    if not pedidas:
        return {}
    try:
        fila_tope = max(int(re.sub(r"[^0-9]", "", c)) for c in pedidas)
    except ValueError:
        return None

    NS = NS_XL
    try:
        with zipfile.ZipFile(str(ruta)) as z:
            nombres = set(z.namelist())
            ruta_hoja, _ = ubicar_hoja_xml(z, hoja)
            if ruta_hoja is None:
                return None

            # 2) recorrer la hoja y cortar al pasar la ultima fila pedida
            compartidas = None
            out = {}
            with z.open(ruta_hoja) as f:
                for evento, el in ET.iterparse(f, events=("end",)):
                    if el.tag == f"{NS}row":
                        try:
                            if int(el.get("r", 0)) > fila_tope:
                                el.clear()
                                break
                        except ValueError:
                            pass
                        el.clear()
                        continue
                    if el.tag != f"{NS}c":
                        continue
                    ref = (el.get("r") or "").upper()
                    if ref not in pedidas:
                        el.clear()
                        continue
                    t = el.get("t")
                    nodo_v = el.find(f"{NS}v")
                    valor = None
                    if t == "inlineStr":
                        nodo_is = el.find(f"{NS}is")
                        if nodo_is is not None:
                            valor = "".join(nodo_is.itertext())
                    elif nodo_v is not None and nodo_v.text is not None:
                        bruto = nodo_v.text
                        if t == "s":                      # texto compartido
                            if compartidas is None:
                                compartidas = []
                                if "xl/sharedStrings.xml" in nombres:
                                    ss = ET.fromstring(z.read("xl/sharedStrings.xml"))
                                    compartidas = ["".join(si.itertext())
                                                   for si in ss.iter(f"{NS}si")]
                            try:
                                valor = compartidas[int(bruto)]
                            except Exception:
                                valor = None
                        elif t == "e":                    # error de formula
                            valor = bruto
                        elif t == "b":
                            valor = bool(int(bruto))
                        elif t == "str":
                            valor = bruto
                        else:                             # numero
                            try:
                                valor = float(bruto)
                            except ValueError:
                                valor = bruto
                    out[ref] = valor
                    el.clear()
                    if len(out) == len(pedidas):
                        break
            return out
    except Exception:
        return None


def diagnosticar_celda(ruta, hoja, celda, log):
    """Cuando una celda no entrega valor, explica POR QUE mirando el XML crudo:
    si esta vacia, si tiene formula sin resultado guardado, si trae texto, si es
    un error, o si es parte de una celda combinada."""
    import zipfile
    import xml.etree.ElementTree as ET

    if not es_zip_excel(ruta):
        log(f"      (no se puede diagnosticar un {Path(ruta).suffix})")
        return
    NS = NS_XL
    ref = celda.upper().replace("$", "")
    try:
        fila_tope = int(re.sub(r"[^0-9]", "", ref))
    except ValueError:
        return
    try:
        with zipfile.ZipFile(str(ruta)) as z:
            ruta_hoja, hojas = ubicar_hoja_xml(z, hoja)
            if ruta_hoja is None:
                log(f"      la hoja '{hoja}' no existe. Hojas: {', '.join(hojas)}")
                return

            encontrada, tiene_formula, formula, tipo, valor = False, False, None, None, None
            with z.open(ruta_hoja) as f:
                for _, el in ET.iterparse(f, events=("end",)):
                    if el.tag == f"{NS}row":
                        try:
                            if int(el.get("r", 0)) > fila_tope:
                                break
                        except ValueError:
                            pass
                        el.clear()
                        continue
                    if el.tag == f"{NS}c" and (el.get("r") or "").upper() == ref:
                        encontrada = True
                        tipo = el.get("t")
                        nf = el.find(f"{NS}f")
                        nv = el.find(f"{NS}v")
                        tiene_formula = nf is not None
                        formula = (nf.text or "")[:70] if nf is not None else None
                        valor = nv.text if nv is not None else None
                        break

            if not encontrada:
                log(f"      {hoja}!{ref} está vacía (la celda no existe en el archivo).")
                # ¿es parte de una celda combinada?
                try:
                    xml = z.read(ruta_hoja).decode("utf-8", "ignore")
                    i = xml.find("<mergeCells")
                    if i > -1:
                        for mr in re.findall(r'ref="([A-Z]+\d+:[A-Z]+\d+)"',
                                             xml[i:xml.find("</mergeCells>", i)]):
                            a, b = mr.split(":")
                            ca, fa = re.match(r"([A-Z]+)(\d+)", a).groups()
                            cb, fb = re.match(r"([A-Z]+)(\d+)", b).groups()
                            cc, fc = re.match(r"([A-Z]+)(\d+)", ref).groups()
                            if (col_letra_a_num(ca) <= col_letra_a_num(cc) <= col_letra_a_num(cb)
                                    and int(fa) <= int(fc) <= int(fb)):
                                log(f"      OJO: está dentro de la celda combinada {mr}; "
                                    f"el valor vive en {a}.")
                                break
                except Exception:
                    pass
                log("      Revisa la hoja y la celda en «Configurar valores...».")
                return

            if tiene_formula and valor is None:
                log(f"      {hoja}!{ref} tiene fórmula pero sin resultado guardado.")
                log(f"      fórmula: ={formula}")
                log("      Pasa cuando el archivo se guardó con cálculo en manual o lo "
                    "guardó otro programa.")
                log("      Solución: ábrelo en Excel, presiona F9 y guárdalo.")
            elif valor is None:
                log(f"      {hoja}!{ref} existe pero no tiene valor guardado.")
            elif tipo == "e":
                log(f"      {hoja}!{ref} tiene un error de fórmula: {valor}")
            else:
                log(f"      {hoja}!{ref} no es un número, es texto: {str(valor)[:50]!r}")
    except Exception as e:
        log(f"      (no se pudo diagnosticar: {e})")


def resolver_hoja(nombres, hoja):
    """Traduce lo pedido al nombre real de la hoja: acepta '#1' (por posicion) y
    tolera tildes y mayusculas. Devuelve None si no calza ninguna."""
    m = re.fullmatch(r"#(\d+)", str(hoja).strip())
    if m:
        i = int(m.group(1)) - 1
        return nombres[i] if 0 <= i < len(nombres) else None
    for n in nombres:
        if n == hoja:
            return n
    for n in nombres:
        if normalizar(n) == normalizar(hoja):
            return n
    return None


def es_significativo(v):
    """Sirve para decidir si una celda 'cuenta' al buscar el ultimo dato de una
    columna: se omiten vacios, ceros y errores de formula (#REF!, #N/D...)."""
    if v is None:
        return False
    if isinstance(v, str):
        t = v.strip()
        return bool(t) and not t.startswith("#")
    if isinstance(v, (int, float)):
        return abs(float(v)) > 0
    return True


leer_columnas_rapido = _excel_xml.leer_columnas_rapido




desescapar_xml = _excel_xml.desescapar_xml


def leer_formulas_rapido(ruta, hoja, columnas, fila_inicio, log):
    """Devuelve {"COL": set(filas_que_TIENEN_formula)}.

    A diferencia de leer_columnas_rapido, que lee el resultado, esto detecta la
    PRESENCIA del nodo <f>, o sea si la celda es una formula o un valor escrito
    a mano. Sirve para saber hasta donde se arrastro una formula.

    Ojo con las formulas compartidas: la primera celda trae la formula completa
    (<f t="shared" ref="L5:L120" si="0">...) y las siguientes solo <f t="shared"
    si="0"/> sin texto. Como aca solo importa que exista un <f>, las dos formas
    cuentan igual.
    """
    import zipfile

    if not es_zip_excel(ruta):
        log(f"    ! {Path(ruta).name}: formato no soportado para leer fórmulas")
        return None
    objetivo = [c.upper() for c in columnas]
    if not objetivo:
        return {}
    alternativas = b"|".join(sorted((c.encode() for c in objetivo),
                                    key=len, reverse=True))
    patron = re.compile(rb'<c r="(' + alternativas + rb')(\d+)"([^>]*?)(?:/>|>(.*?)</c>)',
                        re.S)
    fila_inicio = int(fila_inicio)
    formulas = {c: set() for c in objetivo}
    try:
        with zipfile.ZipFile(str(ruta)) as z:
            ruta_hoja, hojas = ubicar_hoja_xml(z, hoja)
            if ruta_hoja is None:
                log(f"    ! La hoja '{hoja}' no existe. Hojas: {', '.join(hojas)}")
                return None
            with z.open(ruta_hoja) as f:
                cola = b""
                while True:
                    trozo = f.read(1 << 20)
                    if not trozo:
                        break
                    buf = cola + trozo
                    fin = 0
                    for m in patron.finditer(buf):
                        fin = m.end()
                        col, nfila = m.group(1).decode(), int(m.group(2))
                        if nfila < fila_inicio:
                            continue
                        cuerpo = m.group(4) or b""
                        if b"<f" in cuerpo:
                            formulas[col].add(nfila)
                    cola = buf[fin:] if fin else buf[-8192:]
    except Exception as e:
        log(f"    ! No se pudieron leer las fórmulas de {Path(ruta).name}: {e}")
        return None
    return formulas


def armar_tabla(datos, col_clave, cols_valor, log, etiqueta="",
                excluir=("0",), info=None):
    """Convierte {"A":{fila:val}} en {clave_normalizada: (nombre, [valores])}.
    Se salta las filas sin clave, que es lo que aparece como vacios al final.

    excluir: claves normalizadas que NO son empresas y hay que descartar. Por
             defecto el "0", porque ninguna empresa se llama asi y se cuela
             cuando sobran formulas arrastrando ceros.
    info:    dict opcional que se rellena con {"duplicadas", "excluidas",
             "vacias"} para que quien llame decida si eso es un fallo.
    """
    tabla, vacias, duplicadas, excluidas = {}, 0, [], []
    fuera = {normalizar(e) for e in (excluir or ())}
    filas = sorted(datos.get(col_clave, {}))
    for f in filas:
        nombre = datos[col_clave].get(f)
        # Un 0 numerico tampoco es empresa: se descarta igual que el "0" texto.
        if isinstance(nombre, (int, float)) and not isinstance(nombre, bool):
            excluidas.append((f, nombre))
            continue
        if not isinstance(nombre, str) or not nombre.strip():
            vacias += 1
            continue
        clave = normalizar(nombre)
        if clave in fuera:
            excluidas.append((f, nombre.strip()))
            continue
        vals = []
        for c in cols_valor:
            v = datos.get(c, {}).get(f)
            vals.append(float(v) if isinstance(v, (int, float)) else v)
        if clave in tabla:
            duplicadas.append((nombre.strip(), f))
        tabla[clave] = (nombre.strip(), vals)
    if info is not None:
        info.update(duplicadas=duplicadas, excluidas=excluidas, vacias=vacias)
    partes = [f"{len(tabla)} empresa(s)"]
    if vacias:
        partes.append(f"{vacias} fila(s) sin empresa omitidas")
    if excluidas:
        partes.append(f"{len(excluidas)} fila(s) con 0 o vacío descartadas")
    log(f"        {etiqueta}: " + ", ".join(partes))
    if duplicadas:
        log(f"        {etiqueta}: EMPRESAS REPETIDAS -> "
            + ", ".join(f"{n} (fila {f})" for n, f in duplicadas[:8]))
        if len(duplicadas) > 8:
            log(f"        {etiqueta}: ... y {len(duplicadas) - 8} repetición(es) más")
    return tabla




# ---------------------------------------------------------------------------
# Access y resolución de valores configurados
# ---------------------------------------------------------------------------
_dependencias = None

def configurar(dependencias):
    """Inyecta VALORES, CACHE y helpers del punto de entrada sin importarlo."""
    global _dependencias
    _dependencias = dependencias
    globals().update({k: v for k, v in dependencias.items()
                      if not k.startswith("__") and k not in globals()})

def conexion_mdb(ruta):
    import pyodbc
    drivers = [d for d in pyodbc.drivers() if "Microsoft Access Driver" in d]
    if not drivers:
        raise RuntimeError(
            "No hay driver de Access instalado (o es de otra arquitectura que este "
            "Python). Instala 'Microsoft Access Database Engine' de la misma "
            f"arquitectura ({8 * 8 if sys.maxsize > 2**32 else 32} bits)."
        )
    cs = f"DRIVER={{{drivers[0]}}};DBQ={ruta};"
    return pyodbc.connect(cs, autocommit=True)


def listar_tablas_mdb(ruta, log):
    try:
        cn = conexion_mdb(ruta)
        cur = cn.cursor()
        tablas = [r.table_name for r in cur.tables(tableType="TABLE")]
        log(f"    Tablas de {ruta.name}: {', '.join(tablas) if tablas else '(ninguna)'}")
        for t in tablas:
            cols = [r.column_name for r in cur.columns(table=t)]
            log(f"      · {t}: {', '.join(cols)}")
        cn.close()
    except Exception as e:
        log(f"    No se pudo inspeccionar {ruta.name}: {e}")


def obtener_tablas_columnas(ruta):
    """{tabla: [columnas]} de una base Access. {} si no se pudo abrir."""
    out = {}
    cn = None
    try:
        cn = conexion_mdb(ruta)
        cur = cn.cursor()
        for t in [r.table_name for r in cur.tables(tableType="TABLE")]:
            out[t] = [r.column_name for r in cur.columns(table=t)]
    except Exception:
        pass
    finally:
        try:
            if cn is not None:
                cn.close()
        except Exception:
            pass
    return out


def desglose_por_tipo(ruta, tabla, columna, columna_tipo, where, log):
    """Escribe en el log la suma agrupada por tipo. Solo informativo."""
    cn = None
    try:
        cn = conexion_mdb(ruta)
        cur = cn.cursor()
        sql = (f"SELECT [{columna_tipo}], SUM([{columna}]), COUNT(*) "
               f"FROM [{tabla}]")
        if where.strip():
            sql += f" WHERE {where}"
        sql += f" GROUP BY [{columna_tipo}] ORDER BY [{columna_tipo}]"
        cur.execute(sql)
        filas = cur.fetchall()
        if filas:
            log(f"      desglose por {columna_tipo}:")
            for f in filas:
                log(f"        {str(f[0]):<12} {fmt_monto(f[1]):>22}   ({f[2]} filas)")
    except Exception as e:
        log(f"      (no se pudo desglosar por tipo: {e})")
    finally:
        try:
            if cn is not None:
                cn.close()
        except Exception:
            pass


def leer_valor_mdb(ruta, tabla, columna, where, log):
    cn = None
    try:
        cn = conexion_mdb(ruta)
        cur = cn.cursor()
        sql = f"SELECT SUM([{columna}]) FROM [{tabla}]"
        if where.strip():
            sql += f" WHERE {where}"
        cur.execute(sql)
        fila = cur.fetchone()
        return float(fila[0]) if fila and fila[0] is not None else None
    except Exception as e:
        log(f"    ! Error consultando {ruta.name}: {e}")
        return None
    finally:
        try:
            if cn is not None:
                cn.close()
        except Exception:
            pass


def obtener_valor(clave, rutas, log, usar_cache=True):
    """Devuelve (valor, mensaje_error_o_None).
    Si el archivo no cambio desde la ultima lectura y se pide exactamente lo
    mismo, devuelve el valor guardado sin abrir el archivo."""
    spec = VALORES[clave]
    ruta = rutas.get(spec["archivo"])
    if ruta is None:
        return None, f"falta el archivo de origen ({spec['archivo']})"

    huella = huella_spec(spec)
    if usar_cache:
        reg = CACHE.obtener(clave, ruta, huella)
        if reg is not None:
            extra = f", {reg['filas']} fila(s)" if reg.get("filas") else ""
            log(f"    {spec['etiqueta']}")
            log(f"      = {fmt_monto(reg['valor'])}   [ya calculado el "
                f"{reg['leido']}{extra}; {ruta.name} sin cambios, no se abrió]")
            return reg["valor"], None

    if spec["tipo"] == "excel_col":
        if not spec.get("hoja") or not spec.get("columna"):
            listar_hojas(ruta, log)
            return None, f"sin configurar hoja/columna para {clave}"
        filtro = ""
        if spec.get("columna_filtro"):
            filtro = (f"  filtrando {spec['columna_filtro']} = "
                      f"{'/'.join(spec.get('valores_filtro') or ['(nada)'])}")
        log(f"    {spec['etiqueta']}  <-  {spec['hoja']}!{spec['columna']}"
            f"{spec.get('fila_inicio', 2)} hacia abajo{filtro}")
        v, n, desg = leer_columna_excel(
            ruta, spec["hoja"], spec["columna"], int(spec.get("fila_inicio", 2)),
            spec.get("columna_filtro", ""), spec.get("valores_filtro") or [], log)
        if v is None:
            return None, f"no se pudo leer {clave}"
        log(f"      {n} fila(s) sumada(s)")
        if spec.get("columna_filtro") and len(desg) > 1:
            buscados = {normalizar(x) for x in (spec.get("valores_filtro") or [])}
            log(f"      valores presentes en la columna {spec['columna_filtro']}:")
            for k in sorted(desg):
                marca = "<--" if k in buscados else "   "
                log(f"        {marca} {k:<14} {fmt_monto(desg[k]):>22}")
        CACHE.poner(clave, ruta, huella, v, n)
        return v, None

    if spec["tipo"] == "excel_etiqueta":
        if not (spec.get("hoja") and spec.get("columna_etiqueta")
                and spec.get("texto_fila") and spec.get("columna_valor")):
            listar_hojas(ruta, log)
            return None, f"sin configurar la búsqueda por rótulo para {clave}"
        log(f"    {spec['etiqueta']}  <-  {spec['hoja']}: fila donde "
            f"{spec['columna_etiqueta']} = '{spec['texto_fila']}', "
            f"valor en {spec['columna_valor']}")
        v = leer_valor_por_etiqueta(
            ruta, spec["hoja"], spec["columna_etiqueta"], spec["texto_fila"],
            spec["columna_valor"], int(spec.get("fila_inicio", 1)), log)
        if v is not None:
            log(f"      = {fmt_monto(v)}")
            CACHE.poner(clave, ruta, huella, v)
        return v, None if v is not None else f"no se pudo leer {clave}"

    if spec["tipo"] == "excel":
        if not spec.get("hoja") or not spec.get("celda"):
            listar_hojas(ruta, log)
            return None, f"sin configurar hoja/celda para {clave}"
        log(f"    {spec['etiqueta']}  <-  {spec['hoja']}!{spec['celda']}")
        v = leer_valor_excel(ruta, spec["hoja"], spec["celda"], log)
        if v is not None:
            log(f"      = {fmt_monto(v)}")
            CACHE.poner(clave, ruta, huella, v)
        return v, None if v is not None else f"no se pudo leer {clave}"

    if spec["tipo"] == "mdb":
        if not spec.get("tabla") or not spec.get("columna"):
            listar_tablas_mdb(ruta, log)
            return None, f"sin configurar tabla/columna para {clave}"
        wh = spec.get("where", "")
        log(f"    {spec['etiqueta']}  <-  SUM([{spec['columna']}]) de [{spec['tabla']}]"
            + (f" WHERE {wh}" if wh.strip() else ""))
        v = leer_valor_mdb(ruta, spec["tabla"], spec["columna"], wh, log)
        if spec.get("columna_tipo"):
            desglose_por_tipo(ruta, spec["tabla"], spec["columna"],
                              spec["columna_tipo"], wh, log)
        if v is not None:
            CACHE.poner(clave, ruta, huella, v)
        return v, None if v is not None else f"no se pudo leer {clave}"

    return None, f"tipo de origen desconocido en {clave}"
